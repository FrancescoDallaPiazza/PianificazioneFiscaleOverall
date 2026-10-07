"""Compila il modello Nota-spese-trasferte.xlsx da un JSON di trasferte e ripete i controlli.

Uso:  python <percorso>/compila_nota.py dati.json uscita.xlsx
      python <percorso>/compila_nota.py --test

dati.json:
{
  "mese": "2026-11",
  "confermata_il": "2026-12-02",
  "veicolo": "Fiat Tipo", "targa": "GH123KL", "tariffa_aci": 0.52,
  "fonte_tariffa": "Tabelle ACI 2026, 15.000 km",
  "sede": "Villafranca di Verona, Via ...", "casa": "Verona, Via ...",
  "righe": [
    {"data": "2026-11-12", "evento": "12/11 09:00 Rossi S.r.l. - sopralluogo",
     "cliente": "Rossi S.r.l.", "commessa": "FatturaLab 2026/123",
     "motivo": "Sopralluogo per aggiornamento DVR", "destinazione": "Vicenza",
     "indirizzo": "Via Roma 10, Vicenza",
     "confermata": true, "auto": "propria",
     "km_da_sede": 150, "km_da_casa": 130,
     "pedaggi": 10.4, "parcheggi": 0, "vitto": 18, "alloggio": 0, "taxi_treno_bus": 0,
     "tracciabile": "S",
     "allegati": ["percorso Maps", "ricevuta casello Vicenza Est 12/11 5,20", "mail Rossi 10/11"]}
  ]
}

- confermata: true solo se Tazio ha confermato che la trasferta c'è stata e che tutte le spese
  della riga le ha pagate lui (non Telepass o carta aziendale, non offerte dal cliente).
- auto: "propria" oppure "aziendale" / "nessuna" (niente km).
- km = il minore fra km_da_sede e km_da_casa; la partenza è quella del minore.
- Più giorni o più clienti nello stesso viaggio: una riga per data e cliente; i km del giro solo
  sulla prima riga del viaggio, le altre con km 0.

Stampa un rapporto JSON e aggiunge un foglio protetto «Preparazione» per l'approvatore.
"""
import json
import re
import shutil
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote_plus

import openpyxl

TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "Nota-spese-trasferte.xlsx"
PRIMA, ULTIMA = 16, 40          # righe della tabella nel foglio «Nota spese»
MESE_MINIMO = "2026-10"         # primo mese del regolamento trasferte (da confermare con Adami)
SEDE = {"villafranca", "villafranca di verona"}
FRAZIONI = {"alpo", "caluri", "dossobuono", "pizzoletta", "quaderni", "rizza", "rosegaferro"}
RESIDENZA = "verona"
TETTO_VITTO_ALLOGGIO = 180.76   # al giorno, art. 95 c. 3 TUIR
MOTIVI_GENERICI = {"varie", "visita", "visita clienti", "clienti", "lavoro", "riunione", "trasferta"}
SPESE = ("pedaggi", "parcheggi", "vitto", "alloggio", "taxi_treno_bus")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\((vr|verona)\)|\bvr\b", " ", s)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()


def cliente_chiave(s):
    return re.sub(r"(srl|spa|snc|sas|srls)$", "", re.sub(r"[^a-z0-9]", "", norm(s)))


def si(v):
    return str(v).strip().upper() in ("S", "SI", "SÌ", "TRUE", "1", "Y", "YES")


def num(r, k):
    return float(r.get(k) or 0)


def nel_comune_sede(comune):
    c = norm(comune)
    return c in SEDE or any(f in c.split() for f in FRAZIONI) or c.endswith("di villafranca")


def maps(a, b):
    return f"https://www.google.com/maps/dir/?api=1&origin={quote_plus(a)}&destination={quote_plus(b)}"


def controlla(dati):
    if dati["mese"] < MESE_MINIMO:
        raise SystemExit("mese non gestito dal regolamento")
    anno, mese = map(int, dati["mese"].split("-"))
    ok, scartate, avvisi = [], [], []
    if norm(dati.get("veicolo")) in ("", "da compilare") or norm(dati.get("targa")) in ("", "da compilare"):
        raise SystemExit("veicolo o targa mancanti nell'intestazione")
    visti, ricevute = set(), defaultdict(list)
    vitto_alloggio_giorno = defaultdict(float)
    for i, r in enumerate(dati["righe"], 1):
        d = date.fromisoformat(r["data"])
        dest = (r.get("destinazione") or "").strip()
        auto = norm(r.get("auto"))
        kms, kmc = r.get("km_da_sede"), r.get("km_da_casa")
        valori = [num(r, k) for k in SPESE] + [float(x) for x in (kms, kmc) if x not in (None, "")]
        allegati = r.get("allegati") or []
        motivo = None
        if not si(r.get("confermata")):
            motivo = "trasferta o spese non confermate da Tazio"
        elif (d.year, d.month) != (anno, mese):
            motivo = "data fuori dal mese della nota"
        elif not r.get("cliente") or not r.get("commessa") or not dest or not r.get("motivo"):
            motivo = "manca cliente, commessa, motivo o destinazione"
        elif norm(r["motivo"]) in MOTIVI_GENERICI:
            motivo = "motivo generico: descrivere l'attività svolta"
        elif nel_comune_sede(dest):
            motivo = "destinazione nel comune della sede (Villafranca di Verona o frazione)"
        elif auto not in ("propria", "aziendale", "nessuna"):
            motivo = "indicare l'auto usata: propria, aziendale o nessuna"
        elif any(v < 0 for v in valori):
            motivo = "valori negativi"
        elif auto == "propria" and (kms in (None, "") or kmc in (None, "")):
            motivo = "servono i km sia dalla sede sia da casa"
        elif auto != "propria" and max(float(kms or 0), float(kmc or 0)) > 0:
            motivo = "km non rimborsabili: auto aziendale o nessuna auto"
        elif (d, cliente_chiave(r["cliente"])) in visti:
            motivo = "duplicato: stessa data e stesso cliente"
        elif not allegati:
            motivo = "nessun allegato"
        elif num(r, "vitto") + num(r, "alloggio") + num(r, "taxi_treno_bus") > 0 and not si(r.get("tracciabile")):
            # il foglio scarterebbe l'intera riga, km compresi: meglio togliere la spesa in contanti
            motivo = "vitto/alloggio/taxi non pagati con mezzo tracciabile: togliere la spesa e rilanciare"
        if motivo:
            scartate.append({"riga": i, "data": r["data"], "cliente": r.get("cliente"), "motivo": motivo})
            continue
        visti.add((d, cliente_chiave(r["cliente"])))
        if auto == "propria":
            kms, kmc = float(kms), float(kmc)
            r["_km"], r["_partenza"] = (kms, "sede") if kms <= kmc else (kmc, "casa")
        else:
            r["_km"], r["_partenza"] = 0.0, "sede"
        for a in allegati:
            ricevute[norm(a)].append(i)
        c = norm(dest)
        if "villafranca" in c:
            avvisi.append(f"riga {i}: «{dest}» non è Villafranca di Verona? Verificare il comune")
        if c == RESIDENZA:
            avvisi.append(f"riga {i}: trasferta a Verona (residenza): motivo per esteso e prova della visita")
        if r["_km"] == 0 and num(r, "pedaggi") > 0:
            avvisi.append(f"riga {i}: pedaggi senza km")
        vitto_alloggio_giorno[d] += num(r, "vitto") + num(r, "alloggio")
        ok.append(r)
    for a, righe in ricevute.items():
        if len(righe) > 1 and "percorso" not in a:
            avvisi.append(f"allegato «{a}» presente in più righe {righe}: possibile doppio rimborso")
    for d, v in sorted(vitto_alloggio_giorno.items()):
        if v > TETTO_VITTO_ALLOGGIO:
            avvisi.append(f"{d.isoformat()}: vitto+alloggio {v:.2f} € oltre {TETTO_VITTO_ALLOGGIO} €: serve autorizzazione scritta")
    if len(ok) > ULTIMA - PRIMA + 1:
        raise SystemExit(f"troppe righe ({len(ok)}): il modello ne contiene {ULTIMA - PRIMA + 1}")
    tariffa = float(dati["tariffa_aci"])
    totale = sum(round(r["_km"] * tariffa, 2) + sum(num(r, k) for k in SPESE) for r in ok)
    return ok, scartate, avvisi, round(totale, 2)


def compila(dati, uscita):
    ok, scartate, avvisi, totale = controlla(dati)
    ok.sort(key=lambda r: r["data"])
    shutil.copy(TEMPLATE, uscita)
    wb = openpyxl.load_workbook(uscita)
    ws = wb["Nota spese"]
    anno, mese = map(int, dati["mese"].split("-"))
    sede = dati.get("sede") or "Villafranca di Verona"
    casa = dati.get("casa") or "Verona"
    ws["C6"] = date(anno, mese, 1)
    for cella, chiave in (("C9", "veicolo"), ("C10", "targa"), ("C11", "tariffa_aci"), ("C12", "fonte_tariffa")):
        if dati.get(chiave) not in (None, ""):
            ws[cella] = dati[chiave]
    for row in range(PRIMA, ULTIMA + 1):   # svuota esempio e celle da compilare; H, P, Q restano
        for col in "ABCDEFGIJKLMNO":
            ws[f"{col}{row}"] = None
    for row, r in zip(range(PRIMA, ULTIMA + 1), ok):
        partenza = "Villafranca di Verona" if r["_partenza"] == "sede" else "Verona"
        valori = {"A": date.fromisoformat(r["data"]), "B": r["cliente"], "C": r["commessa"],
                  "D": r["motivo"], "E": partenza, "F": r["destinazione"], "G": r["_km"],
                  "I": num(r, "pedaggi"), "J": num(r, "parcheggi"), "K": num(r, "vitto"),
                  "L": num(r, "alloggio"), "M": num(r, "taxi_treno_bus"),
                  "N": "S" if si(r.get("tracciabile")) else "N", "O": len(r["allegati"])}
        for col, v in valori.items():
            ws[f"{col}{row}"] = v
        ws[f"A{row}"].number_format = "dd/mm/yyyy"

    # traccia per l'approvatore
    p = wb.create_sheet("Preparazione")
    p.append(["NOTA PREPARATA CON ASSISTENTE DA GOOGLE CALENDAR — NON APPROVATA"])
    p.append([f"Generata il {datetime.now():%d/%m/%Y %H:%M}; dati confermati da Tazio Pradella il "
              f"{dati.get('confermata_il') or '[data]'}"])
    p.append([])
    p.append(["Data", "Evento di calendario", "Cliente", "Km da sede", "Km da casa", "Km usati",
              "Percorso sede", "Percorso casa", "Allegati"])
    for r in ok:
        dest = r.get("indirizzo") or r["destinazione"]
        p.append([r["data"], r.get("evento", ""), r["cliente"], r.get("km_da_sede"), r.get("km_da_casa"),
                  r["_km"], maps(sede, dest), maps(casa, dest), "; ".join(r["allegati"])])
    p.append([])
    p.append(["Righe scartate"])
    for s in scartate:
        p.append([s["data"], "", s["cliente"], s["motivo"]])
    p.append([])
    p.append(["Avvisi per l'approvatore"])
    for a in avvisi + ["Doppioni con note di mesi diversi: da verificare a cura dell'approvatore."]:
        p.append([a])
    p.protection.sheet = True
    wb.save(uscita)
    return {"file": str(uscita), "righe_ok": len(ok), "totale_rimborsabile": totale,
            "scartate": scartate, "avvisi": avvisi}


def _test():
    import tempfile
    base = dict(cliente="Rossi S.r.l.", commessa="FL 1", motivo="Sopralluogo DVR", destinazione="Vicenza",
                confermata=True, auto="propria", km_da_sede=100, km_da_casa=120, pedaggi=10,
                parcheggi=0, vitto=20, alloggio=0, taxi_treno_bus=0, tracciabile="S",
                allegati=["percorso", "prova visita"])
    dati = {"mese": "2026-11", "veicolo": "Fiat Tipo", "targa": "GH123KL", "tariffa_aci": 0.5, "righe": [
        dict(base, data="2026-11-03"),                                     # OK: 50+10+20 = 80
        dict(base, data="2026-11-03", cliente="Rossi srl"),                # duplicato (grafia diversa)
        dict(base, data="2026-11-04", destinazione="Dossobuono di Villafranca"),
        dict(base, data="2026-11-05", destinazione="Villafranca di Verona (VR)"),
        dict(base, data="2026-10-30"),                                     # fuori mese
        dict(base, data="2026-11-06", tracciabile="N"),                    # contanti
        dict(base, data="2026-11-07", allegati=[]),                        # senza allegati
        dict(base, data="2026-11-08", confermata=False),                   # non confermata
        dict(base, data="2026-11-09", auto="aziendale"),                   # km con auto aziendale
        dict(base, data="2026-11-11", km_da_sede=-50),                     # negativo
        dict(base, data="2026-11-12", motivo="varie"),                     # generico
        dict(base, data="2026-11-10", cliente="Bianchi", destinazione="Verona (VR)",
             km_da_sede=40, km_da_casa=8, vitto=90, alloggio=100),         # OK: 4+10+190, oltre tetto
        dict(base, data="2026-11-13", cliente="Gamma", destinazione="Villafranca Padovana"),  # OK + avviso
    ]}
    out = Path(tempfile.mkdtemp()) / "t.xlsx"
    rep = compila(dati, out)
    assert rep["righe_ok"] == 3, rep
    assert len(rep["scartate"]) == 10, rep
    assert rep["totale_rimborsabile"] == 80 + 204 + 80, rep
    av = " ".join(rep["avvisi"])
    assert "oltre" in av and "residenza" in av and "Padovana" in av, rep
    wb = openpyxl.load_workbook(out)
    ws = wb["Nota spese"]
    assert ws["B16"].value == "Rossi S.r.l." and ws["E17"].value == "Verona" and ws["G17"].value == 8
    assert ws["N16"].value == "S" and ws["B19"].value is None
    assert str(ws["Q16"].value).startswith("=IF(") and ws.protection.sheet
    assert wb["Preparazione"].protection.sheet
    try:
        compila(dict(dati, mese="2026-09"), out)
        raise AssertionError("mese minimo non applicato")
    except SystemExit as e:
        assert str(e) == "mese non gestito dal regolamento"
    print("test ok")


if __name__ == "__main__":
    if sys.argv[1:] == ["--test"]:
        _test()
    else:
        dati = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
        print(json.dumps(compila(dati, Path(sys.argv[2])), ensure_ascii=False, indent=2))
